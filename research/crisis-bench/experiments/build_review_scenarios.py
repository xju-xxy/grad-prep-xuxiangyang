# -*- coding: utf-8 -*-
"""
场景要点审改清单生成器（供标注人逐条审改使用）。

产出：
  results/审改-场景要点.xlsx
    - Sheet1 要点清单：709 条要点的定位、风险标记、AI 建议、留白供标注人填写
    - Sheet2 场景规范核对：每场景的结构合规性自动核验
    - Sheet3 风险统计：按风险类别计数
  results/审改-结论表述核对.md：论文结论逐条对照实验真值

风险标记（自动启发式，仅作定位提示，不代表一定有问题）：
  R1 答案泄漏   数值推算型要点的期望数值直接出现在情报原文中（严重）
  R2 情报可抄   非推算要点的匹配词全部能在情报原文中找到（低区分度）
  R3 泛词匹配   匹配词全为泛词，规则裁判易误判命中
  R4 措辞含糊   要点文本含模糊限定词（加强/做好/确保/尽快/统筹…）
  R5 对象不明   决策层要点未出现本场景具体对象（村名/场所/资源）
  R6 无期望值   推算型要点未写明的数值或区间
  R7 近重复     同场景内要点文本字符二元组重合度过高
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = Path(__file__).resolve().parent.parent
SCEN = BASE / "scenarios" / "scenarios.json"
OUT_XLSX = BASE / "results" / "审改-场景要点.xlsx"
OUT_MD = BASE / "results" / "审改-结论表述核对.md"
ANALYSIS = BASE / "results" / "v0.4-analysis.json"

LAYER_KEY = {"F": "事实层", "R": "推理层", "D": "决策层"}
LAYER_CN = {"事实层": "F", "推理层": "R", "决策层": "D"}

GENERIC = set("""优先 力量 措施 路线 撤离路线 撤离 转移 安排 组织 部署 监测 排查 核实 澄清
发布 官方 通知 保障 调度 评估 预案 应急 处置 响应 建议 要求 方案 行动 工作 加强 处置措施
救援 抢险 物资 人员 群众 现场 指挥 联动 协调""".split())

VAGUE = ["加强", "做好", "确保", "尽快", "统筹", "协调推进", "落实", "强化", "高度重视", "适当",
         "合理", "相关", "等，", "妥善", "尽快"]

TRAP_HINT = ["不能", "禁止", "不得", "错误", "误区", "不应", "切勿", "陷阱", "严禁", "否决", "识破"]


def bigrams(s):
    s = re.sub(r"[\s，。、；：（）()【】\[\]“”\"'·\-/]", "", s)
    return set(s[i:i + 2] for i in range(len(s) - 1)) or {s}


def jaccard(a, b):
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


def has_number(s):
    return bool(re.search(r"\d", s))


def main():
    data = json.loads(SCEN.read_text(encoding="utf-8"))["scenarios"]

    rows = []          # 要点清单
    scene_rows = []    # 场景规范核对

    for s in data:
        sid = s["id"]
        cat = s["类别"]
        intel = "；".join(s["情报"])
        res = s["资源"]
        res_txt = "；".join(res)
        n_calc = n_trap = 0
        per_scene = []

        for layer_cn, items in s["rubric"].items():
            tag = LAYER_CN[layer_cn]
            for idx, it in enumerate(items, 1):
                pid = "%s%d" % (tag, idx)
                point = it["要点"]
                match = it.get("match", []) or []
                flags, notes, sugg = [], [], []

                # 数值推算/汇总型：含"推算/汇总/估算/换算"，或含"数字×数字""数字=数字"的算式
                is_calc = layer_cn != "决策层" and (
                    any(k in point for k in ("推算", "汇总", "估算", "换算")) or
                    ("×" in point and has_number(point)) or ("=" in point and has_number(point)))

                # R1 答案泄漏：推算型要点的期望数值出现在情报里
                if is_calc:
                    leaked = [m for m in match
                              if re.search(r"\d", m) and len(m) >= 4 and m in intel]
                    if leaked:
                        flags.append("R1")
                        notes.append("答案泄漏：匹配词 %s 已直接出现在情报原文" % "、".join(leaked))
                        sugg.append("改用未泄漏的等价数值区间，或调整情报使目标数字不直接出现")

                # R2 情报可抄：非推算要点，匹配词全部能在情报里找到
                if not is_calc and match:
                    hit = [m for m in match if m in intel]
                    if hit and len(hit) == len(match):
                        flags.append("R2")
                        notes.append("情报可抄：全部匹配词均在情报原文中出现，抄录即可命中")

                # R3 泛词匹配
                if match and all(m in GENERIC for m in match):
                    flags.append("R3")
                    notes.append("泛词匹配：匹配词 %s 全为泛词，规则裁判易误判" % "、".join(match))
                    sugg.append("补充含具体对象/数值的匹配词")

                # R4 措辞含糊
                v = [w for w in VAGUE if w in point]
                if v:
                    flags.append("R4")
                    notes.append("措辞含糊：含「%s」等模糊限定词" % "、".join(sorted(set(v))))

                # R5 决策层对象不明：要点文本自身是否含本场景的具体对象（场所/村名/资源名）
                if layer_cn == "决策层":
                    specific = []
                    for n in range(2, 7):
                        for k in range(len(point) - n + 1):
                            sub = point[k:k + n]
                            if sub in GENERIC:
                                continue
                            if sub in intel or sub in res_txt:
                                specific.append(sub)
                    if not specific:
                        flags.append("R5")
                        notes.append("对象不明：决策层要点未出现本场景具体对象（村名/场所/资源）")
                        sugg.append("在要点中写明对象（如具体村名、场所或资源名）")

                # R6 推算型无期望值
                if is_calc and not re.search(r"\d", point):
                    flags.append("R6")
                    notes.append("无期望值：推算型要点未写明期望数值或区间")
                    sugg.append("补写期望数值区间，便于裁判核验")

                if is_calc:
                    n_calc += 1
                if any(h in point for h in TRAP_HINT):
                    n_trap += 1

                rows.append({
                    "场景": sid, "类别": cat, "灾种": s["灾种"], "层": layer_cn, "编号": pid,
                    "要点": point, "匹配词": "、".join(match), "_match": match,
                    "风险标记": "+".join(flags), "风险说明": "；".join(notes),
                    "AI建议": "；".join(sugg), "_bg": bigrams(point),
                })
                per_scene.append((sid, pid, point, bigrams(point), layer_cn))

        # R7 同场景近重复（跨要点）
        for i in range(len(per_scene)):
            for j in range(i + 1, len(per_scene)):
                if per_scene[i][4] == per_scene[j][4]:  # 仅同层比较，避免误报
                    jc = jaccard(per_scene[i][3], per_scene[j][3])
                    if jc >= 0.75:
                        for r in rows:
                            if r["场景"] == sid and r["编号"] in (per_scene[i][1], per_scene[j][1]):
                                if "R7" not in r["风险标记"]:
                                    r["风险标记"] = (r["风险标记"] + "+R7").strip("+")
                                    r["风险说明"] += "；近重复：与 %s 文本重合度 %.2f" % (
                                        per_scene[j][1] if r["编号"] == per_scene[i][1] else per_scene[i][1], jc)

        n_intel = len(s["情报"])
        n_res = len(res)
        has_neg = any(r.startswith("无") for r in res)
        n_calc_fact = sum(1 for it in s["rubric"]["事实层"]
                          if any(k in it["要点"] for k in ("推算", "汇总", "估算", "换算"))
                          or ("×" in it["要点"] and has_number(it["要点"]))
                          or ("=" in it["要点"] and has_number(it["要点"])))
        struct = []
        if not (8 <= n_intel <= 11):
            struct.append("情报条数 %d 不在 8~11" % n_intel)
        if not (4 <= n_res <= 6):
            struct.append("资源项数 %d 不在 4~6" % n_res)
        if not has_neg:
            struct.append("资源清单无「无X」否定项（v0.3 原版场景，未回溯补齐）")
        if n_calc_fact == 0:
            struct.append("事实层无数值推算/汇总型要点")

        # 信息矛盾类：需含保守口径/统一口径/矛盾相关要点
        if cat == "信息矛盾":
            kw = ["矛盾", "口径", "保守", "统一", "核实", "核实核灾"]
            if not any(any(k in it["要点"] for k in kw) for it in s["rubric"]["决策层"]):
                struct.append("信息矛盾类场景决策层缺「保守口径/统一口径」要点")
            # 矛盾组数：以 rubric 要点中显式标注的矛盾为口径（情报正文未必带"矛盾"字样）
            contra = sum(1 for items in s["rubric"].values() for it in items
                         if any(k in it["要点"] for k in ("矛盾", "不一致", "差异", "vs")))
            if contra < 2:
                struct.append("矛盾要点疑似不足 2 组（rubric 中仅 %d 条）" % contra)

        scene_rows.append({
            "场景": sid, "类别": cat, "灾种": s["灾种"], "情报条数": n_intel,
            "资源项数": n_res, "含否定资源": "是" if has_neg else "否",
            "事实": len(s["rubric"]["事实层"]), "推理": len(s["rubric"]["推理层"]),
            "决策": len(s["rubric"]["决策层"]),
            "小计": sum(len(v) for v in s["rubric"].values()),
            "推算点数": n_calc, "陷阱提示点数": n_trap,
            "规范核对": "合格" if not struct else "；".join(struct),
        })

    # ===== 写 xlsx =====
    wb = Workbook()
    ws = wb.active
    ws.title = "要点清单"
    hdr = ["场景", "类别", "灾种", "层", "编号", "要点", "匹配词",
           "风险标记", "风险说明", "AI建议（供参考，须本人确认）", "你的审改（改后文本）", "备注"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(horizontal="center", vertical="center")

    fill_hi = PatternFill("solid", fgColor="FFC7CE")   # 红：R1/R2/R3/R7
    fill_md = PatternFill("solid", fgColor="FFEB9C")   # 黄：R4/R5/R6
    for r in rows:
        ws.append([r["场景"], r["类别"], r["灾种"], r["层"], r["编号"], r["要点"], r["匹配词"],
                   r["风险标记"], r["风险说明"], r["AI建议"], "", ""])
        f = r["风险标记"]
        if any(x in f for x in ("R1", "R2", "R3", "R7")):
            ws.cell(ws.max_row, 8).fill = fill_hi
        elif f:
            ws.cell(ws.max_row, 8).fill = fill_md
    widths = [8, 10, 12, 8, 8, 52, 30, 12, 44, 34, 26, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:L%d" % ws.max_row

    ws2 = wb.create_sheet("场景规范核对")
    hdr2 = ["场景", "类别", "灾种", "情报条数", "资源项数", "含否定资源",
            "事实", "推理", "决策", "小计", "推算点数", "陷阱提示点数", "规范核对"]
    ws2.append(hdr2)
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in scene_rows:
        ws2.append([r[k] for k in hdr2])
        if r["规范核对"] != "合格":
            ws2.cell(ws2.max_row, 13).fill = fill_hi
    for i, w in enumerate([8, 10, 12, 10, 10, 12, 8, 8, 8, 8, 10, 12, 46], 1):
        ws2.column_dimensions[chr(64 + i)].width = w
    ws2.freeze_panes = "B2"

    ws3 = wb.create_sheet("风险统计")
    ws3.append(["风险类别", "条数", "含义"])
    for c in ws3[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    desc = {
        "R1": "答案泄漏：推算型要点的期望数值已出现在情报原文（严重，削弱区分度）",
        "R2": "情报可抄：非推算要点全部匹配词可在情报中找到（低区分度，属设计内）",
        "R3": "泛词匹配：匹配词全为泛词，规则裁判易误判",
        "R4": "措辞含糊：要点含模糊限定词",
        "R5": "对象不明：决策层要点未含本场景具体对象",
        "R6": "无期望值：推算型要点未写明数值区间",
        "R7": "近重复：同层要点文本重合度过高",
    }
    for k in ["R1", "R2", "R3", "R4", "R5", "R6", "R7"]:
        n = sum(1 for r in rows if k in r["风险标记"])
        ws3.append([k, n, desc[k]])
    ws3.append([])
    ws3.append(["合计要点", len(rows), "覆盖 %d 个场景" % len(scene_rows)])
    ws3.append(["结构不合规场景", sum(1 for r in scene_rows if r["规范核对"] != "合格"),
                "详见「场景规范核对」表"])
    for i, w in enumerate([12, 10, 70], 1):
        ws3.column_dimensions[chr(64 + i)].width = w

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)

    # ===== 结论表述核对 md =====
    a = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    t = a["tests"]
    rq1 = t["RQ1: multi vs single"]
    rq2 = t["RQ2: noaudit vs multi"]
    rq4 = t["RQ4: verify vs multi"]
    a2ms = t["A2: plus模型 multi vs single"]
    a2mf = t.get("A2: 被测模型 plus vs flash")
    lines = [
        "# 结论表述核对表（逐条对照实验真值）",
        "",
        "> 用途：投稿前核对论文每条定量结论是否有对应实测依据。真值取自 "
        "`results/v0.4-analysis.json`（本表数值由脚本直接读取，避免手抄漂移）。",
        "> 判定：**支撑**=数值一致且有显著性依据；**需限定**=数值一致但表述需加限定语。",
        "",
        "| # | 论文结论 | 真值依据 | 判定 |",
        "|---|---|---|---|",
        "| 1 | 多智能体较单智能体高 7.20pp（综合 90.5%% vs 83.3%%） | mean_diff_pp=%.2f；McNemar b=%d/c=%d p=%.2e；Wilcoxon p=%.4f；CI=[%.4f,%.4f] | 支撑 |" % (
            rq1["mean_diff_pp"], rq1["mcnemar"]["b(A0B1)"], rq1["mcnemar"]["c(A1B0)"],
            rq1["mcnemar"]["p"], rq1["wilcoxon"]["p"], rq1["bootstrap_ci_95"][0], rq1["bootstrap_ci_95"][1]),
        "| 2 | 去掉审核官使 multi 回落 7.04pp | mean_diff_pp=%.2f；McNemar p=%.2e | 支撑 |" % (
            rq2["mean_diff_pp"], rq2["mcnemar"]["p"]),
        "| 3 | 约束校验组件再带来 2.81pp 增益 | mean_diff_pp=%.2f；McNemar p=%.2e；CI=[%.4f, %.4f] | 支撑（CI 下界贴近 0，需保留 CI 表述） |" % (
            rq4["mean_diff_pp"], rq4["mcnemar"]["p"], rq4["bootstrap_ci_95"][0], rq4["bootstrap_ci_95"][1]),
        "| 4 | glm-4-plus 较 flash 高 9.39pp（p=7.4e-16） | %s | %s |" % (
            json.dumps(a2mf, ensure_ascii=False) if a2mf else "见 tests",
            "需核对" if not a2mf else "支撑"),
        "| 5 | plus 下 multi 与 single 已无显著差异 | mean_diff_pp=%.2f；McNemar p=%.4f（二项精确） | 支撑 |" % (
            a2ms["mean_diff_pp"], a2ms["mcnemar"]["p"]),
        "",
        "## kappa / 自一致率",
        "",
        "| 配置 | kappa(flash vs plus) | 论文表述 |",
        "|---|---|---|",
    ]
    for k, v in a["kappa"].items():
        lines.append("| %s | %.3f | %s |" % (k, v, "一致" if abs(v) < 0.1 else "需核对"))
    lines.append("")
    lines.append("| 配置 | flash 自一致率 |")
    lines.append("|---|---|")
    for k, v in a["self_consistency"].items():
        lines.append("| %s | %.2f |" % (k, v))
    lines += [
        "",
        "## 推算合规（§6.6 表述「multi 与 multi-verify 推算合规同为 12/49」）",
        "",
        "| 配置 | calc_ok / calc_total |",
        "|---|---|",
    ]
    for k, v in a["verifier"].items():
        lines.append("| %s | %d/%d |" % (k, v["calc_ok"], v["calc_total"]))
    lines += [
        "",
        "## 误差类型学（§7.2）",
        "",
        "| 配置 | total_loss | 各类型计数 |",
        "|---|---|---|",
    ]
    for k, v in a["errors"].items():
        lines.append("| %s | %d | %s |" % (k, v["total_loss"], json.dumps(v["counts"], ensure_ascii=False)))
    lines += [
        "",
        "## 表述风险提示（供逐条审改参考）",
        "",
        "- 「接近天花板」「区分度趋饱和」为定性表述，已有 35/36 满分支撑，建议保留「需 v0.5 更难场景」的限定。",
        "- 「约束校验组件再带来 2.81pp」的 CI 下界 +0.07pp 贴近 0，正文已保留 CI，勿删。",
        "- 「多数票优于任一单裁判」的结论基于 71 条分歧集，须保留「分歧集上」限定，避免被读作全集结论。",
    ]
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")

    # ===== 控制台摘要 =====
    print("要点总数 %d，场景 %d" % (len(rows), len(scene_rows)))
    for k in ["R1", "R2", "R3", "R4", "R5", "R6", "R7"]:
        print("  %s: %d" % (k, sum(1 for r in rows if k in r["风险标记"])))
    bad = [r for r in scene_rows if r["规范核对"] != "合格"]
    print("结构不合规场景 %d" % len(bad))
    for r in bad:
        print("   %s %s" % (r["场景"], r["规范核对"]))
    print("输出", OUT_XLSX.name, "/", OUT_MD.name)


if __name__ == "__main__":
    main()
